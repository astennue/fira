#!/usr/bin/env python3
"""Build FIRA Data Dictionary workbook (algorithm-ready, code-verified enums)."""
import sys, os, datetime

SKILL_DIR = "/home/z/my-project/skills/xlsx"
for p in (SKILL_DIR, os.path.join(SKILL_DIR, "templates"), "/home/z/my-project/scripts"):
    if p not in sys.path:
        sys.path.insert(0, p)

from base import *  # design tokens + helpers (single source of truth)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from fira_dd_data1 import ORG_TABLES, APPLICANT_TABLES
from fira_dd_data2 import (JOB_TABLES, ATS_TABLES, ENDORSEMENT_TABLES,
                           AI_TABLES, SUPPORT_TABLES, CMS_TABLES)
from fira_dd_data3 import (STATUS_MACHINE, PIPELINE, FORMULAS, FEATURES,
                           QUALITY_NOTES, RELATIONSHIPS, FLOW_EDGES, GROUPS,
                           COLUMN_LEGEND, ROLE_LEGEND, README_ABOUT)

OUT = "/home/z/my-project/download/FIRA-Data-Dictionary.xlsx"
TODAY = datetime.date.today().isoformat()

WRAP_LT = Alignment(horizontal="left", vertical="top", wrap_text=True)
WRAP_CTR = Alignment(horizontal="center", vertical="top", wrap_text=True)

ENDORSEMENT_FLOW = [
    ("create endorsement", "FIRA staff / super_admin / international_agency / local_agency", "(new)", "pending_fira_review", "Application created by applicant sits in FIRA review queue"),
    ("fira_approve", "FIRA staff / super_admin / international_agency", "pending_fira_review", "pending_employer_review", "Application.status -> pending_employer_review"),
    ("fira_reject", "FIRA staff / super_admin / international_agency", "pending_fira_review", "fira_rejected", "Application.status -> fira_rejected (recoverable to applied/screening)"),
    ("employer_accept", "employer / FIRA staff", "pending_employer_review", "employer_accepted", "Application.status -> employer_accepted (then offered -> hired -> deployed)"),
    ("employer_decline", "employer / FIRA staff", "pending_employer_review", "employer_declined", "Application.status -> employer_declined (terminal)"),
]

def section_row(ws, row, col_start, col_end, text):
    # NOTE: no merged cells (auto_fit_columns cannot handle MergedCell);
    # span the fill across the range instead.
    for col in range(col_start, col_end + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = PatternFill("solid", fgColor=SECONDARY)
    c = ws.cell(row=row, column=col_start, value=text)
    c.font = Font(name=FONT_NAME, size=12, bold=HEADER_BOLD, color=PRIMARY)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 24

def write_table(ws, start_row, headers, rows, center_cols=()):
    """Write header + data rows starting at start_row, col B. Returns next free row."""
    last_col = len(headers) + 1
    for i, h in enumerate(headers, start=2):
        ws.cell(row=start_row, column=i, value=h)
    style_header_row(ws, start_row, 2, last_col)
    r = start_row + 1
    for idx, row in enumerate(rows):
        for i, v in enumerate(row, start=2):
            ws.cell(row=r, column=i, value=v)
        style_data_row(ws, r, 2, last_col, idx)
        for col in range(2, last_col + 1):
            ws.cell(row=r, column=col).alignment = WRAP_CTR if col in center_cols else WRAP_LT
        ws.row_dimensions[r].height = None  # auto-fit wrapped text
        r += 1
    return r

def caption(ws, row, col, text):
    c = ws.cell(row=row, column=col, value=text)
    c.font = font_caption()
    c.alignment = Alignment(horizontal="left", vertical="center")

# ============================================================
# Workbook
# ============================================================
wb = Workbook()
wb.properties.creator = "Z.ai"

# ---------- README ----------
ws = wb.active
ws.title = "README"
ws.sheet_properties.tabColor = PRIMARY
setup_sheet(ws, title="FIRA Data Dictionary - Job Matching & Ranking Platform", last_col=5)
r = 4
section_row(ws, r, 2, 5, "About This Dictionary"); r += 1
r = write_table(ws, r, ["Item", "Description"],
                [(k, f"{v}") for k, v in README_ABOUT] + [("Generated", TODAY)])
r += 1
section_row(ws, r, 2, 5, "Column Legend (applies to every dictionary sheet)"); r += 1
r = write_table(ws, r, ["Column", "Definition"], COLUMN_LEGEND)
r += 1
section_row(ws, r, 2, 5, "Algorithm Role Legend (values of the 'Algorithm Role' column)"); r += 1
r = write_table(ws, r, ["Tag", "Definition"], ROLE_LEGEND)
r += 1
section_row(ws, r, 2, 5, "Entity Inventory"); r += 1

grp_data = [
    ("01 User & Org", ORG_TABLES), ("02 Applicant", APPLICANT_TABLES),
    ("03 Job Orders", JOB_TABLES), ("04 Applications & ATS", ATS_TABLES),
    ("05 Endorsement", ENDORSEMENT_TABLES), ("06 AI & Matching", AI_TABLES),
    ("07 System & Support", SUPPORT_TABLES), ("08 CMS", CMS_TABLES),
]
inv_rows = []
for (sheet_name, _tables, focus), (_n, tables) in zip(GROUPS, grp_data):
    inv_rows.append((sheet_name, ", ".join(t["name"] for t in tables),
                     sum(len(t["fields"]) for t in tables), focus))
r = write_table(ws, r, ["Sheet", "Tables Covered", "Field Count", "Data Focus"], inv_rows, center_cols=(4,))
r += 1
caption(ws, r, 2, "All enumerations are code-verified against the repository (src/app/api/*, src/lib/status.ts, src/components/*, python-ai/*) as of " + TODAY + ".")
ws.freeze_panes = "B5"
auto_fit_columns(ws, min_width=10, max_width=80, header_row=5, data_start_row=6)
ws.column_dimensions["C"].width = 78

# ---------- Algorithm Mapping ----------
ws = wb.create_sheet("Algorithm Mapping")
ws.sheet_properties.tabColor = ACCENT_POSITIVE
setup_sheet(ws, title="Algorithm Mapping - Candidate <-> Job Matching & Ranking", last_col=7)
r = 4
section_row(ws, r, 2, 7, "A. Matching Pipeline (end-to-end, code-verified)"); r += 1
r = write_table(ws, r, ["Step", "Stage", "Component / Code", "Inputs (Table.Field)", "Output", "Description"], PIPELINE, center_cols=(2,))
r += 1
section_row(ws, r, 2, 7, "B. Scoring Formula Variants"); r += 1
r = write_table(ws, r, ["Variant", "Formula", "Score Range", "When Used"], FORMULAS, center_cols=(4,))
r += 1
section_row(ws, r, 2, 7, "C. Feature-to-Field Mapping (what feeds the score)"); r += 1
r = write_table(ws, r, ["Feature", "Source (Table.Field)", "Scale", "Weight", "Notes"], FEATURES, center_cols=(4, 5))
r += 1
section_row(ws, r, 2, 7, "D. Data Quality & Alignment Notes (for algorithm tuning)"); r += 1
r = write_table(ws, r, ["Topic", "Current Behavior", "Why It Matters", "Recommendation"], QUALITY_NOTES)
ws.freeze_panes = "B5"
auto_fit_columns(ws, min_width=9, max_width=60, header_row=5, data_start_row=6)
for col, w in {"C": 34, "D": 30, "E": 34, "F": 34, "G": 46}.items():
    ws.column_dimensions[col].width = w

# ---------- Relationship Map ----------
ws = wb.create_sheet("Relationship Map")
setup_sheet(ws, title="Relationship Map - Entities & Matching Data Flow", last_col=7)
r = 4
section_row(ws, r, 2, 7, "Entity Relationships (Prisma schema)"); r += 1
r = write_table(ws, r, ["From Table", "Relationship", "To Table", "Key / FK", "Cardinality", "Description"], RELATIONSHIPS, center_cols=(6,))
r += 1
section_row(ws, r, 2, 7, "Matching Data Flow (field -> algorithm -> output)"); r += 1
r = write_table(ws, r, ["Source", "", "Destination"], FLOW_EDGES, center_cols=(3,))
r += 1
caption(ws, r, 2, "Soft reference = linked by convention in application code, without a database-level foreign key constraint.")
ws.freeze_panes = "B5"
auto_fit_columns(ws, min_width=9, max_width=60, header_row=5, data_start_row=6)
ws.column_dimensions["E"].width = 40
ws.column_dimensions["G"].width = 52

# ---------- Dictionary sheets ----------
DICT_HEADERS = ["Table", "Field", "Type", "Required", "Description",
                "Allowed Values / Format", "Populated By", "Related Table (FK)", "Algorithm Role"]

def build_dict_sheet(sheet_name, tables, extra_blocks=None, focus=""):
    ws = wb.create_sheet(sheet_name)
    setup_sheet(ws, title=f"{sheet_name} - Data Dictionary", last_col=10)
    r = 4
    for t in tables:
        section_row(ws, r, 2, 10, f"Table: {t['name']}")
        r += 1
        cap = ws.cell(row=r, column=2, value="Purpose: " + t["purpose"])
        cap.font = font_caption()
        cap.alignment = WRAP_LT
        ws.row_dimensions[r].height = None
        r += 1
        rows = [(t["name"],) + f for f in t["fields"]]
        r = write_table(ws, r, DICT_HEADERS, rows, center_cols=(4, 5))
        r += 1
    for block_title, headers, rows, center in (extra_blocks or []):
        section_row(ws, r, 2, 10, block_title)
        r += 1
        r = write_table(ws, r, headers, rows, center_cols=center)
        r += 1
    caption(ws, r, 2, "Source: prisma/schema.prisma + code audit of API routes, UI constants and the python-ai service. Generated " + TODAY + ".")
    ws.freeze_panes = "C5"
    auto_fit_columns(ws, min_width=8, max_width=46, header_row=4, data_start_row=5)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    return ws

build_dict_sheet("01 User & Org", ORG_TABLES)
build_dict_sheet("02 Applicant", APPLICANT_TABLES)
build_dict_sheet("03 Job Orders", JOB_TABLES)
build_dict_sheet("04 Applications & ATS", ATS_TABLES, extra_blocks=[(
    "Application Status State Machine (code-verified: src/lib/status.ts)",
    ["Status", "Category", "Meaning", "Who Can Set (Roles)", "Valid Next Statuses", "Terminal?"],
    STATUS_MACHINE, (3, 7))])
build_dict_sheet("05 Endorsement", ENDORSEMENT_TABLES, extra_blocks=[(
    "Endorsement Status Flow (code-verified: src/app/api/endorsements/route.ts)",
    ["Action", "Performed By", "From Status", "To Status", "Effect on Application.status"],
    ENDORSEMENT_FLOW, (4, 5))])
build_dict_sheet("06 AI & Matching", AI_TABLES)
build_dict_sheet("07 System & Support", SUPPORT_TABLES)
build_dict_sheet("08 CMS", CMS_TABLES)

os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)

# ---------- summary ----------
total_fields = sum(len(t["fields"]) for _, tables in grp_data for t in tables)
total_tables = sum(len(tables) for _, tables in grp_data)
print(f"Saved: {OUT}")
print(f"Tables documented: {total_tables} | Fields documented: {total_fields}")
print("Sheets:", wb.sheetnames)
